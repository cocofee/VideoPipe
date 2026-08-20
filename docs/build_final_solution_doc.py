from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "cyclerace_multisource_finish_review_final.docx"
REFERENCE_IMAGE = Path(
    r"C:\Users\Administrator\AppData\Local\Temp\orca-paste-1787218185405-cabd9695-959b-4543-a8ac-9e43bbec0406.png"
)

FONT = "Microsoft YaHei"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(95, 103, 112)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B8C2CC", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa) -> None:
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=11, color=None, bold=None, italic=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_paragraph(p, before=0, after=6, line=1.25, align=None):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        p.alignment = align


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True, color=INK)
        set_run_font(p.add_run(text[len(bold_prefix):]), color=INK)
    else:
        set_run_font(p.add_run(text), color=INK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=4)
    set_run_font(p.add_run(text), color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, after=4)
    set_run_font(p.add_run(text), color=INK)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, before=2, after=8, line=1.05)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.08)
    run = p.add_run(text)
    set_run_font(run, size=9.5, color=INK, name="Consolas")
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT)
    p_pr.append(shd)
    return p


def add_callout(doc, label, text, fill="E8EEF5"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    set_table_borders(table, color="9EB6CE", size="8")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    style_paragraph(p, before=2, after=2, line=1.15)
    set_run_font(p.add_run(label + "  "), color=DARK_BLUE, bold=True)
    set_run_font(p.add_run(text), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    set_table_borders(table)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        style_paragraph(p, after=0, line=1.05, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run_font(p.add_run(header), size=9.5, color=DARK_BLUE, bold=True)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            if row_idx % 2 == 1:
                set_cell_shading(cells[idx], "FAFBFC")
            p = cells[idx].paragraphs[0]
            style_paragraph(p, after=0, line=1.08, align=WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            set_run_font(p.add_run(str(value)), size=9.5, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        color, size, before, after = BLUE, 16, 18, 10
    elif level == 2:
        color, size, before, after = BLUE, 13, 14, 7
    else:
        color, size, before, after = DARK_BLUE, 12, 10, 5
    style_paragraph(p, before=before, after=after, line=1.1)
    set_run_font(p.add_run(text), size=size, color=color, bold=True)
    return p


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    for name, size, color in (("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK_BLUE)):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(11)


def configure_page(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(header, after=0, line=1.0)
    set_run_font(header.add_run("CycleRace 多源终点判读系统 | 最终方案"), size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(footer, after=0, line=1.0)
    set_run_font(footer.add_run("内部技术方案 | 2026-08-20"), size=8.5, color=MUTED)


def build_doc():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)

    p = doc.add_paragraph()
    style_paragraph(p, before=12, after=4, line=1.0)
    set_run_font(p.add_run("最终技术方案"), size=10, color=BLUE, bold=True)
    p = doc.add_paragraph()
    style_paragraph(p, after=4, line=1.0)
    set_run_font(p.add_run("CycleRace 多源终点判读系统"), size=25, color=INK, bold=True)
    p = doc.add_paragraph()
    style_paragraph(p, after=12, line=1.1)
    set_run_font(p.add_run("FinishLynx 式单窗口判读 + 单/双摄像头 + 澳亚特线阵 + CycleRace 官方计时"), size=12, color=MUTED)
    add_callout(doc, "最终决策", "第一版先实现 CycleRace + 单普通摄像头 + VideoPipe；双摄像头和澳亚特线阵分别作为后续增强，CycleRace 始终是唯一官方成绩来源。", fill="E8EEF5")

    if REFERENCE_IMAGE.exists():
        p = doc.add_paragraph()
        style_paragraph(p, before=2, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run().add_picture(str(REFERENCE_IMAGE), width=Inches(6.15))
        cap = doc.add_paragraph()
        style_paragraph(cap, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run_font(cap.add_run("图 1  FinishLynx 单窗口布局参考：赛事信息、结果表、线阵图像和普通视频同时可见"), size=9, color=MUTED, italic=True)

    add_heading(doc, "1. 决策摘要", 1)
    add_body(doc, "本方案采用 FinishLynx 的单窗口判读思路，但不复制其代码、内部协议或商业实现。系统由 CycleRace、VideoPipe 和澳亚特软件组成。")
    add_body(doc, "核心边界：CycleRace 是唯一官方成绩来源；VideoPipe 负责视频检索、人工判读、证据保存和修正建议；澳亚特继续负责高速线阵采集。")
    add_bullet(doc, "普通摄像头用于连续视频证据和号码确认。")
    add_bullet(doc, "高速线阵用于精确判断冲线顺序和冲线位置。")
    add_bullet(doc, "YOLO、OCR 都不是第一版的必要条件。")
    add_bullet(doc, "原始视频、RGB 和人工观察记录必须保留。")

    add_heading(doc, "2. 系统组成和职责", 1)
    add_table(doc, ["组件", "官方职责", "与其他组件的关系"], [
        ("CycleRace", "芯片计时、通过顺序、正式名次和成绩", "向 VideoPipe 发送 passage 事件；人工确认后处理正式修正"),
        ("VideoPipe", "视频检索、人工判读、证据保存、修正建议", "整合普通视频、双摄像头和澳亚特 RGB"),
        ("澳亚特软件", "高速线阵相机采集、生成 RGB", "保持原软件稳定运行，由 VideoPipe 只读接入"),
    ], [1800, 3300, 4260])
    add_body(doc, "当前 VideoPipe 仓库已经具备多路视频管线基础：RTSP/UDP/文件输入、异步录像、按 channel_index 管理多路数据和应用层回调。需要新增比赛事件模型、源 PTS 时间索引、统一时间轴和判读界面。")

    add_heading(doc, "3. 支持的工作模式", 1)
    add_table(doc, ["模式", "设备", "主要用途", "优先级"], [
        ("基础模式", "CycleRace + 1 台普通摄像头", "按芯片时间找视频、确认号码、保存证据", "第一版"),
        ("双摄模式", "CycleRace + 2 台普通摄像头", "两个角度确认号码、减少遮挡", "第二阶段"),
        ("精确模式", "CycleRace + 澳亚特线阵", "判断贴身冲线顺序和精确位置", "线阵接入后"),
        ("完整模式", "CycleRace + 双摄 + 澳亚特", "芯片、顺序、号码三源互证", "最终推荐"),
    ], [1600, 2650, 3650, 1460])
    add_heading(doc, "3.1 基础模式", 2)
    add_body(doc, "普通摄像头连续录像，CycleRace 每收到一个芯片通过事件，就在 VideoPipe 的视频时间线上增加标记。操作员点击标记即可跳到目标附近。这一模式不依赖 YOLO、OCR 或高速相机。")
    add_heading(doc, "3.2 双摄模式", 2)
    add_code(doc, "摄像头 A：终点正面或正侧面，重点看车头和正面号码\n摄像头 B：终点反侧面或背面，解决遮挡、重叠和背号不可见")
    add_body(doc, "两路视频分别保存，不预先拼接。VideoPipe 使用一个主时间线同步控制两路回放，每路保留独立的 camera_id、源 PTS 和时间偏移。")
    add_heading(doc, "3.3 精确和完整模式", 2)
    add_body(doc, "澳亚特线阵回答‘谁先通过、前轮前缘在哪一列、多个运动员的相对顺序是什么’；普通摄像头回答‘这个通过对象的号码是多少、是否存在芯片漏读’。")

    add_heading(doc, "4. 网络拓扑", 1)
    add_code(doc, "CycleRace 电脑 ─┐\nVideoPipe 电脑 ─┼── 专用千兆交换机\n澳亚特电脑   ──┤\n摄像头 A     ──┤\n摄像头 B     ──┘")
    add_bullet(doc, "使用固定 IP，设备网和办公网隔离。")
    add_bullet(doc, "某台电脑原网口已连接设备时，增加 USB 千兆网卡。")
    add_bullet(doc, "关键事件使用 TCP、HTTP 或 WebSocket；UDP 只做状态广播。")
    add_bullet(doc, "录像和证据使用独立 SSD，并设置磁盘空间告警。")

    add_heading(doc, "5. FinishLynx 风格主界面", 1)
    add_code(doc, "┌─────────────────────────────────────────────────────┐\n│ 菜单、工具栏、赛事状态、设备状态、待判数量             │\n├───────────────────┬─────────────────────────────────┤\n│ 赛事和设备信息     │ CycleRace 通过记录/判读结果表     │\n├───────────────────┼─────────────────────────────────┤\n│ 澳亚特线阵图像     │ 普通摄像头 A 或双摄像头画面       │\n│ 垂直判读线         │ 播放、暂停、逐帧、慢放            │\n├───────────────────┴─────────────────────────────────┤\n│ 统一时间线、芯片标记、线阵位置、视频位置               │\n├─────────────────────────────────────────────────────┤\n│ 上一人 下一人 绑定号码 标记未知 保存证据 建立修正建议   │\n└─────────────────────────────────────────────────────┘")
    add_body(doc, "顶部状态区显示赛事、组别、圈次、CycleRace、澳亚特、摄像头 A/B 的连接状态、同步误差和待判读数量。")
    add_body(doc, "结果表字段：通过序号、芯片号码、芯片时间、线阵判读号码、视频判读号码、圈次、状态和证据。")
    add_body(doc, "结果表的每一行都是一个可定位的判读对象，而不仅是最终成绩表。设备未连接时隐藏对应画面，不影响其他模式运行。")

    add_heading(doc, "6. 统一时间轴和联动", 1)
    add_code(doc, "点击结果表记录\n    ↓\n摄像头 A 跳转\n    ↓\n摄像头 B 跳转\n    ↓\n澳亚特线阵判读线跳转\n\n拖动统一时间线：摄像头 A、摄像头 B、线阵图像同步移动")
    add_body(doc, "第一版体验目标是点击芯片事件后跳到视频前后约 1 秒范围，再由操作员拖动和逐帧确认。正式版本需要源 PTS、统一时钟和现场校准。")
    add_callout(doc, "核心难点", "真正的技术难点不是同时打开两路视频，而是让 CycleRace 时间、普通视频 PTS、澳亚特 device_tick 和 VideoPipe 本地时间可靠地映射到同一条时间轴。", fill="FFF4D6")
    add_body(doc, "必须保存：CycleRace passage_time_ms、普通视频 source_pts_us、澳亚特 device_tick、VideoPipe monotonic_time、消息接收时间、视频文件和帧号、每台设备的 offset。不能把网络到达时间或 frame_index/FPS 简单换算当作拍摄时间。")

    add_heading(doc, "7. 澳亚特 RGB 适配", 1)
    add_table(doc, ["项目", "已确认内容", "实现要求"], [
        ("文件头", "48 字节", "不足一个完整头时等待下一次读取"),
        ("扫描记录", "每条 3080 字节", "只解析完整记录，避免半条记录"),
        ("像素", "1024 个 CCD 像素 × 3 通道", "生成线阵总图和局部放大图"),
        ("标志", "0x01000000 开始；0x02000000 结束", "支持孤立、短片段和文件重置异常"),
    ], [1800, 3100, 4460])
    add_body(doc, "适配器以只读方式监视 RGB 文件增长，为每段保存独立元数据和图像索引，不覆盖原始 RGB，也不依赖 JPG 作为实时接口。")

    add_heading(doc, "8. 现场业务流程", 1)
    add_heading(doc, "8.1 正常通过", 2)
    add_code(doc, "CycleRace 收到芯片 -> 发送 passage 事件 -> VideoPipe 增加时间线标记\n-> 操作员点击记录 -> 普通视频和线阵图像跳转\n-> 人工确认号码和状态 -> 保存判读和证据")
    add_heading(doc, "8.2 芯片漏读", 2)
    add_code(doc, "线阵发现实际有 4 个通过位置，CycleRace 只有 3 条芯片记录\n-> VideoPipe 建立匿名通过槽位\n-> 普通视频确认号码\n-> 生成‘疑似芯片漏读’建议\n-> CycleRace 操作员确认\n-> 正式成绩由 CycleRace 处理")
    add_callout(doc, "现场要求", "普通摄像头必须连续录像或使用循环缓冲，不能只在收到芯片后才录像，否则没有芯片的运动员将没有视频定位点。", fill="FFF4D6")

    add_heading(doc, "9. 数据对象和通信", 1)
    add_code(doc, "RaceEvent: event_id, race_id, stage_id, group_id, lap\nPassageEvent: passage_id, sequence, chip_id, bib, passage_time_ms, lap, source\nVideoFrameIndex: camera_id, source_pts_us, frame_index, file_path, keyframe_offset\nAytSegment: file_id, segment_index, start_tick, end_tick, start_line, end_line\nReviewObservation: observation_id, passage_id, manual_bib, line_position, video_pts_us, status\nEvidenceItem: image_path, video_clip_path, source_hash, created_at")
    add_body(doc, "CycleRace -> VideoPipe 的消息必须支持确认、去重、断线重连、历史补发、本地持久化和协议版本。VideoPipe -> CycleRace 只发送修正建议，不直接写正式成绩。")

    add_heading(doc, "10. 证据和审计", 1)
    add_bullet(doc, "保存比赛、组别、圈次和 CycleRace 原始事件。")
    add_bullet(doc, "保存视频文件、源 PTS、摄像头 A/B 帧位置和 RGB 片段位置。")
    add_bullet(doc, "保存截图、视频片段、人工判读结果、操作员和时间。")
    add_bullet(doc, "原始录像和 RGB 不覆盖，修改采用追加记录。")
    add_bullet(doc, "VideoPipe 重启后恢复事件和待判队列。")

    add_heading(doc, "11. 开发计划", 1)
    add_table(doc, ["阶段", "核心任务", "验收重点"], [
        ("0 接口确认", "CycleRace 字段、摄像头 PTS、RGB 共享、固定 IP、时间同步", "接口字段表、拓扑图、校准记录"),
        ("1 事件接口", "接收、确认、去重、重连、补发、本地持久化", "1000 条事件无丢失无重复"),
        ("2 单摄 MVP", "RTSP 连续录像、源 PTS、时间索引、点击跳转、逐帧回放", "一个窗口完成视频检索"),
        ("3 判读界面", "事件窗口、结果表、号码绑定、未知对象、待判队列", "新事件不打断当前判读"),
        ("4 双摄", "两路独立录像、offset、同步回放、断流重连、双路证据", "可显示并校准同步偏差"),
        ("5 证据建议", "截图、片段、审计、哈希、修正建议", "每条建议可回到原始证据"),
        ("6 澳亚特", "RGB 增量解析、分段、判读线、tick 映射、三源联动", "下一段不覆盖上一段"),
        ("7 现场验收", "单人、集团、漏读、遮挡、断流、重启、长时间运行", "记录误差、耗时和恢复情况"),
        ("8 可选 AI", "YOLO/OCR 候选和运动员跟踪", "仅作为辅助，不改变权威边界"),
    ], [1700, 5000, 2660])

    add_heading(doc, "12. 核心验收标准", 1)
    for item in [
        "1000 条 passage 事件无丢失、无重复。",
        "断线恢复后可以补齐事件。",
        "单摄点击事件可以定位到目标前后约 1 秒范围。",
        "双摄可以独立校准并显示同步偏差。",
        "原始视频、RGB、截图和人工记录全部可追溯。",
        "芯片漏读时仍可以建立匿名通过观察。",
        "澳亚特分段完成后及时进入待判队列。",
        "新事件不强制打断当前判读。",
        "VideoPipe 不绕过人工确认直接修改正式成绩。",
    ]:
        add_number(doc, item)

    add_heading(doc, "13. 风险与暂不解决事项", 1)
    add_body(doc, "主要风险包括摄像头源 PTS 不稳定、三台电脑时钟漂移、普通摄像头遮挡、澳亚特实时共享或 SDK 权限未确认、长时间双路录像的存储压力，以及 CycleRace/VideoPipe 接口字段尚未冻结。")
    add_code(doc, "暂不做：反编译 FinishLynx；复制其内部代码；依赖 OCR 才能判读；依赖 YOLO 才能保存证据；VideoPipe 自动修改正式成绩；只保存 JPG 而不保存原始证据")

    add_heading(doc, "14. 最终结论", 1)
    add_callout(doc, "最终形态", "CycleRace 负责官方计时和正式成绩；VideoPipe 负责统一判读界面、时间线和证据；普通摄像头负责连续视频和号码确认；双摄像头减少遮挡；澳亚特线阵负责精确判断冲线顺序。", fill="E8EEF5")
    add_body(doc, "第一版从 CycleRace + 单普通摄像头 + VideoPipe 开始；随后增加双摄像头，最后接入澳亚特线阵，形成适合自行车比赛的多源终点判读系统。")

    add_heading(doc, "参考资料", 1)
    for item in [
        "FinishLynx 官方帮助包：third_party_reference/finishlynx_13.10/help_extracted/Content/OnlineManual/",
        "EventWindow.htm、ImageZoneExplanations.htm、ResultsZoneExplanations.htm",
        "IdentiLynx.htm、MultiUserMode.htm、MYLAPS.htm",
        "澳亚特 RGB 解析基础：scripts/parse_ayt_rgb.py",
        "用户提供的 FinishLynx 界面截图（图 1）",
    ]:
        add_bullet(doc, item)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
